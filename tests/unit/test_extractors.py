"""Extractores con muestras sintéticas. El OCR real se marca pendiente si falta Tesseract."""

from pathlib import Path

import docx
import pytest

from docia_e.adapters.extraction.inspect import inspect_file
from docia_e.adapters.extraction.service import extract_document, ocr_status
from docia_e.domain.errors import DomainError


LIMITS = {
    "max_upload_bytes": 25 * 1024 * 1024,
    "max_upload_mib": 25,
    "max_pdf_pages": 100,
    "max_image_pixels": 30_000_000,
    "docx_max_uncompressed_bytes": 200 * 1024 * 1024,
    "docx_max_entries": 5000,
}


def _pdf(pages: list[str]) -> bytes:
    objects: list[bytes] = []

    def add(body: str) -> int:
        objects.append(body.encode("latin-1"))
        return len(objects)

    font_id = add("<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>")
    content_ids = []
    for text in pages:
        escaped = _pdf_text(text)
        stream = f"BT /F1 12 Tf 40 100 Td ({escaped}) Tj ET"
        content_ids.append(add(f"<< /Length {len(stream.encode('latin-1'))} >>\nstream\n{stream}\nendstream"))
    page_ids = []
    for content_id in content_ids:
        page_ids.append(
            add(
                f"<< /Type /Page /Parent PAGES 0 R /MediaBox [0 0 595 842] "
                f"/Contents {content_id} 0 R /Resources << /Font << /F1 {font_id} 0 R >> >> >>"
            )
        )
    kids = " ".join(f"{page_id} 0 R" for page_id in page_ids)
    pages_id = add(f"<< /Type /Pages /Count {len(page_ids)} /Kids [{kids}] >>")
    add(f"<< /Type /Catalog /Pages {pages_id} 0 R >>")
    built = []
    for index, body in enumerate(objects, start=1):
        if b"Parent PAGES 0 R" in body:
            body = body.replace(b"Parent PAGES 0 R", f"Parent {pages_id} 0 R".encode("ascii"))
        built.append(body)
    header = b"%PDF-1.4\n"
    chunks = [header]
    offsets = [0]
    cursor = len(header)
    for index, body in enumerate(built, start=1):
        offsets.append(cursor)
        piece = f"{index} 0 obj\n".encode("ascii") + body + b"\nendobj\n"
        chunks.append(piece)
        cursor += len(piece)
    xref = cursor
    out = b"".join(chunks)
    xref_lines = ["xref", f"0 {len(built) + 1}", "0000000000 65535 f "]
    for offset in offsets[1:]:
        xref_lines.append(f"{offset:010d} 00000 n ")
    trailer = (
        "\n".join(xref_lines)
        + f"\ntrailer << /Size {len(built) + 1} /Root {len(built)} 0 R >>\nstartxref\n{xref}\n%%EOF"
    )
    return out + trailer.encode("ascii")


def _pdf_text(text: str) -> str:
    encoded = text.encode("latin-1")
    chars = []
    for byte in encoded:
        if byte in (0x28, 0x29, 0x5C):
            chars.append("\\" + chr(byte))
        elif 32 <= byte < 127:
            chars.append(chr(byte))
        else:
            chars.append(f"\\{byte:03o}")
    return "".join(chars)


def test_pdf_digital_conserva_orden_y_no_habilita_t(tmp_path: Path):
    path = tmp_path / "jornada.pdf"
    path.write_bytes(_pdf(["No hay orientación. RA4 CE2.1. 40 %.", "Segunda página."]))
    inspected = inspect_file(path, "jornada.pdf", LIMITS)
    assert inspected["kind"] == "pdf"
    result = extract_document(path, "pdf", {"ocr_languages": "spa", "max_pdf_pages": 100, "max_image_pixels": 30_000_000})
    texts = [unit["text"] for unit in result["units"]]
    assert texts[0].startswith("No hay")
    assert "40 %" in texts[0]
    assert "Segunda" in texts[1]
    assert [unit["locator"]["page"] for unit in result["units"]] == [1, 2]
    assert "ready" not in result["run_status"] or result["run_status"] in {"succeeded", "partial", "failed"}


def test_docx_conserva_parrafo_tabla_parrafo_sin_pagina(tmp_path: Path):
    path = tmp_path / "guia.docx"
    document = docx.Document()
    document.add_paragraph("Antes de la tabla, no se pierde este párrafo.")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "RA4"
    table.rows[0].cells[1].text = "40 %"
    document.add_paragraph("Después de la tabla.")
    document.save(path)
    inspect_file(path, "guia.docx", LIMITS)
    result = extract_document(path, "docx", {})
    texts = [unit["text"] for unit in result["units"]]
    assert texts[0].startswith("Antes")
    assert "RA4" in texts
    assert "40 %" in texts
    assert texts[-1].startswith("Después")
    assert texts.index("Antes de la tabla, no se pierde este párrafo.") < texts.index("RA4")
    assert all(unit["locator"]["page"] is None for unit in result["units"])
    assert all(unit["locator"]["type"] == "docx_block" for unit in result["units"])


def test_extension_falsa_y_zip_que_no_es_docx(tmp_path: Path):
    fake = tmp_path / "foto.pdf"
    fake.write_bytes(b"\x89PNG\r\n\x1a\n" + b"0" * 20)
    with pytest.raises(DomainError) as caught:
        inspect_file(fake, "foto.pdf", LIMITS)
    assert caught.value.code == "EXTENSION_MISMATCH"
    package = tmp_path / "malo.docx"
    import zipfile

    with zipfile.ZipFile(package, "w") as archive:
        archive.writestr("readme.txt", "no es un documento")
    with pytest.raises(DomainError) as caught:
        inspect_file(package, "malo.docx", LIMITS)
    assert caught.value.code == "UNSUPPORTED_FORMAT"


def test_ocr_no_se_da_por_aprobado_si_falta_tesseract(tmp_path: Path):
    status = ocr_status()
    if status["available"] and "spa" in status["languages"]:
        pytest.skip("Tesseract está instalado; la ausencia se prueba en otro entorno")
    path = tmp_path / "vacio.pdf"
    path.write_bytes(_pdf([" "] * 1))
    result = extract_document(path, "pdf", {"ocr_languages": "spa", "max_pdf_pages": 100, "max_image_pixels": 30_000_000})
    codes = {issue["code"] for issue in result["issues"]}
    assert "OCR_UNAVAILABLE" in codes or "EMPTY_TEXT" in codes
    assert result["run_status"] != "succeeded" or "OCR_UNAVAILABLE" in codes
